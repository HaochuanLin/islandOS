"""Tests for the Philippines DIB workbook importer."""
import csv
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook

from experiments.import_philippines_dib_workbook import import_workbook


def make_workbook(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "02 - Screened Mini-grids"
    ws.append(["Grid Number", "Grid Name", "Longitude [°]", "Latitude [°]", "Daily Schedule"])
    ws.append([1, "Test Island", 121.1, 13.2, 24])

    ws = wb.create_sheet("03 - Monthly Power Demand")
    ws.append(["Grid Number", "Grid Name", "Maximum Load [MW]", "Minimum Load [MW]",
               "January 2022 Load [MW]", None, None])
    ws.append([None, None, None, None, "Maximum", "Minimum", "Average"])
    ws.append([1, "Test Island", 0.3, 0.1, 0.25, 0.12, 0.2])

    for name, unit in (("05 - Monthly GHI", "kW/m2"),
                       ("06 - Monthly Wind Speed", "m/s")):
        ws = wb.create_sheet(name)
        ws.append(["Grid Number", "Grid Name", f"Resource [{unit}]"])
        ws.append([None, None, *["Jan", "Feb", "Mar", "Apr", "May", "Jun",
                                  "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]])
        ws.append([1, "Test Island", *[0.1] * 12])

    for name, first_header in (("07 - Diesel Capacities", "Dependable Capacity [MW]"),
                               ("08 - Diesel Accumulated Hours", "Accumulated Operating Hours [h]")):
        ws = wb.create_sheet(name)
        ws.append(["Grid Number", "Grid Name", first_header])
        ws.append([None, None, "Diesel 1"])
        ws.append([1, "Test Island", 0.25 if "Capacities" in name else 1234.0])
    wb.save(path)


class PhilippinesDibImporterTests(unittest.TestCase):
    def test_imports_tidy_files_and_warns_about_monthly_resolution(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            workbook = root / "sample.xlsx"
            output = root / "out"
            make_workbook(workbook)

            manifest = import_workbook(workbook, output)

            self.assertEqual(manifest["site_count"], 1)
            self.assertEqual(manifest["row_counts"]["monthly_power_demand.csv"], 1)
            self.assertEqual(manifest["row_counts"]["monthly_ghi.csv"], 12)
            with (output / "monthly_power_demand.csv").open(encoding="utf-8", newline="") as handle:
                row = next(csv.DictReader(handle))
            self.assertEqual(row["month"], "2022-01")
            self.assertEqual(row["temporal_resolution"], "monthly_summary")
            with (output / "monthly_wind_speed.csv").open(encoding="utf-8", newline="") as handle:
                resource = next(csv.DictReader(handle))
            self.assertEqual(resource["month_of_year"], "01")
            self.assertEqual(resource["wind_speed_unit"], "m/s")
            self.assertTrue(any("not hourly observations" in note for note in manifest["data_limitations"]))
            self.assertTrue((output / "diesel_assets.csv").is_file())
            self.assertTrue((output / "manifest.json").is_file())

    def test_rejects_workbook_with_missing_required_sheets(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            workbook = root / "incomplete.xlsx"
            Workbook().save(workbook)
            with self.assertRaisesRegex(ValueError, "missing required sheets"):
                import_workbook(workbook, root / "out")


if __name__ == "__main__":
    unittest.main()
