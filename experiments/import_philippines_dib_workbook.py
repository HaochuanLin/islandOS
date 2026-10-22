"""Normalize the Philippines DIB workbook into tidy, auditable CSV files.

The importer preserves source units and monthly resolution. It does not turn
monthly summaries into hourly observations or verify the workbook's original
licence/provenance.
"""
from __future__ import annotations

import argparse
import csv
import json
import math
import re
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

SITES = "02 - Screened Mini-grids"
DEMAND = "03 - Monthly Power Demand"
GHI = "05 - Monthly GHI"
WIND = "06 - Monthly Wind Speed"
CAPACITY = "07 - Diesel Capacities"
HOURS = "08 - Diesel Accumulated Hours"
REQUIRED_SHEETS = (SITES, DEMAND, GHI, WIND, CAPACITY, HOURS)
MONTHS = ("Jan", "Feb", "Mar", "Apr", "May", "Jun",
          "Jul", "Aug", "Sep", "Oct", "Nov", "Dec")


def _numeric(value: Any, field: str, row: int) -> float:
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        raise ValueError(f"{field} must be numeric at worksheet row {row}")
    result = float(value)
    if not math.isfinite(result):
        raise ValueError(f"{field} must be finite at worksheet row {row}")
    return result


def _grid_id(value: Any, row: int) -> int:
    number = _numeric(value, "Grid Number", row)
    if number < 0 or not number.is_integer():
        raise ValueError(f"Grid Number must be a non-negative integer at worksheet row {row}")
    return int(number)


def _write_csv(path: Path, fields: list[str], rows: list[dict[str, Any]]) -> None:
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        writer.writerows(rows)


def _month_from_header(value: Any, column: int) -> str:
    if not isinstance(value, str):
        raise ValueError(f"Missing monthly power-demand header at column {column}")
    match = re.match(r"^([A-Za-z]+) (\d{4}) Load \[MW\]$", value.strip())
    if not match:
        raise ValueError(f"Unexpected monthly power-demand header at column {column}: {value!r}")
    try:
        return datetime.strptime(f"{match.group(1)} {match.group(2)}", "%B %Y").strftime("%Y-%m")
    except ValueError as exc:
        raise ValueError(f"Unrecognized month in header: {value!r}") from exc


def import_workbook(input_path: str | Path, output_dir: str | Path) -> dict[str, Any]:
    """Import site, monthly demand/resource, and diesel-asset tables."""
    source, output = Path(input_path), Path(output_dir)
    if not source.is_file():
        raise FileNotFoundError(f"Workbook not found: {source}")
    output.mkdir(parents=True, exist_ok=True)
    workbook = load_workbook(source, read_only=True, data_only=True)
    try:
        missing = [name for name in REQUIRED_SHEETS if name not in workbook.sheetnames]
        if missing:
            raise ValueError("Workbook missing required sheets: " + ", ".join(missing))

        # Site registry.
        ws = workbook[SITES]
        sites, site_ids = [], set()
        for row in range(2, ws.max_row + 1):
            raw_id = ws.cell(row, 1).value
            if raw_id is None:
                continue
            grid_id = _grid_id(raw_id, row)
            name = ws.cell(row, 2).value
            if not isinstance(name, str) or not name.strip():
                raise ValueError(f"Grid Name is blank at {SITES} row {row}")
            if grid_id in site_ids:
                raise ValueError(f"Duplicate Grid Number {grid_id} in {SITES}")
            site_ids.add(grid_id)
            site = {"grid_number": grid_id, "grid_name": name.strip()}
            for column, field in ((3, "longitude_deg"), (4, "latitude_deg"),
                                  (5, "daily_schedule_hours")):
                value = ws.cell(row, column).value
                site[field] = "" if value is None else _numeric(value, field, row)
            sites.append(site)
        if not sites:
            raise ValueError("No grid rows found in screened mini-grids sheet")
        _write_csv(output / "sites.csv",
                   ["grid_number", "grid_name", "longitude_deg", "latitude_deg",
                    "daily_schedule_hours"], sites)

        # Monthly load data: three subcolumns per month.
        ws = workbook[DEMAND]
        demand_rows, demand_ids = [], set()
        for row in range(3, ws.max_row + 1):
            raw_id = ws.cell(row, 1).value
            if raw_id is None:
                continue
            grid_id = _grid_id(raw_id, row)
            demand_ids.add(grid_id)
            name = str(ws.cell(row, 2).value or "").strip()
            for start in range(5, ws.max_column + 1, 3):
                header = ws.cell(1, start).value
                if header is None:
                    continue
                month = _month_from_header(header, start)
                values = [ws.cell(row, start + offset).value for offset in range(3)]
                if all(value is None for value in values):
                    continue
                maximum, minimum, average = [
                    _numeric(value, label, row)
                    for value, label in zip(values, ("Maximum Load", "Minimum Load", "Average Load"))
                ]
                if min(maximum, minimum, average) < 0:
                    raise ValueError(f"Negative load for grid {grid_id}, month {month}")
                tolerance = max(1e-9, abs(maximum) * 1e-9)
                if minimum > average + tolerance or average > maximum + tolerance:
                    raise ValueError(f"Expected minimum <= average <= maximum for grid {grid_id}, month {month}")
                demand_rows.append({
                    "grid_number": grid_id, "grid_name": name, "month": month,
                    "maximum_load_mw": maximum, "minimum_load_mw": minimum,
                    "average_load_mw": average, "temporal_resolution": "monthly_summary",
                })
        if demand_ids != site_ids:
            raise ValueError("Grid IDs in monthly demand do not match the screened grid registry")
        _write_csv(output / "monthly_power_demand.csv",
                   ["grid_number", "grid_name", "month", "maximum_load_mw",
                    "minimum_load_mw", "average_load_mw", "temporal_resolution"], demand_rows)

        # Monthly resource sheets: keep the source's GHI unit verbatim and don't invent a year.
        def read_resource(sheet_name: str, value_field: str, unit_field: str,
                          source_unit: str) -> list[dict[str, Any]]:
            resource_ws = workbook[sheet_name]
            records, found_ids = [], set()
            for row in range(3, resource_ws.max_row + 1):
                raw_id = resource_ws.cell(row, 1).value
                if raw_id is None:
                    continue
                grid_id = _grid_id(raw_id, row)
                found_ids.add(grid_id)
                name = str(resource_ws.cell(row, 2).value or "").strip()
                for col, month_abbr in enumerate(MONTHS, start=3):
                    value = resource_ws.cell(row, col).value
                    if value is None:
                        continue
                    number = _numeric(value, value_field, row)
                    if number < 0:
                        raise ValueError(f"Negative {value_field} for grid {grid_id}, month {month_abbr}")
                    records.append({
                        "grid_number": grid_id, "grid_name": name,
                        "month_of_year": f"{datetime.strptime(month_abbr, '%b').month:02d}",
                        value_field: number, unit_field: source_unit,
                    })
            if found_ids != site_ids:
                raise ValueError(f"Grid IDs in {sheet_name} do not match the screened grid registry")
            return records

        ghi_rows = read_resource(GHI, "ghi_value_as_reported", "ghi_unit_as_reported", "kW/m2")
        wind_rows = read_resource(WIND, "wind_speed_10m_m_s", "wind_speed_unit", "m/s")
        _write_csv(output / "monthly_ghi.csv",
                   ["grid_number", "grid_name", "month_of_year", "ghi_value_as_reported",
                    "ghi_unit_as_reported"], ghi_rows)
        _write_csv(output / "monthly_wind_speed.csv",
                   ["grid_number", "grid_name", "month_of_year", "wind_speed_10m_m_s",
                    "wind_speed_unit"], wind_rows)

        # Pair generator capacity and accumulated hours using their unit labels.
        cap_ws, hour_ws = workbook[CAPACITY], workbook[HOURS]
        cap_cols = {str(cap_ws.cell(2, col).value).strip(): col
                    for col in range(3, cap_ws.max_column + 1)
                    if cap_ws.cell(2, col).value is not None}
        hour_cols = {str(hour_ws.cell(2, col).value).strip(): col
                     for col in range(3, hour_ws.max_column + 1)
                     if hour_ws.cell(2, col).value is not None}
        hours_by_grid, hour_ids = {}, set()
        for row in range(3, hour_ws.max_row + 1):
            raw_id = hour_ws.cell(row, 1).value
            if raw_id is None:
                continue
            grid_id = _grid_id(raw_id, row)
            hour_ids.add(grid_id)
            hours_by_grid[grid_id] = {unit: hour_ws.cell(row, col).value
                                      for unit, col in hour_cols.items()}
        diesel_rows, capacity_ids = [], set()
        for row in range(3, cap_ws.max_row + 1):
            raw_id = cap_ws.cell(row, 1).value
            if raw_id is None:
                continue
            grid_id = _grid_id(raw_id, row)
            capacity_ids.add(grid_id)
            name = str(cap_ws.cell(row, 2).value or "").strip()
            hours = hours_by_grid.get(grid_id, {})
            for unit, col in cap_cols.items():
                capacity, operating_hours = cap_ws.cell(row, col).value, hours.get(unit)
                if capacity is None and operating_hours is None:
                    continue
                if capacity is not None and _numeric(capacity, "Dependable Capacity", row) < 0:
                    raise ValueError(f"Negative capacity for grid {grid_id}, unit {unit}")
                if operating_hours is not None and _numeric(operating_hours, "Operating Hours", row) < 0:
                    raise ValueError(f"Negative operating hours for grid {grid_id}, unit {unit}")
                diesel_rows.append({
                    "grid_number": grid_id, "grid_name": name, "unit_name": unit,
                    "dependable_capacity_mw": "" if capacity is None else float(capacity),
                    "accumulated_operating_hours": "" if operating_hours is None else float(operating_hours),
                })
        if capacity_ids != site_ids or hour_ids != site_ids:
            raise ValueError("Grid IDs in diesel sheets do not match the screened grid registry")
        _write_csv(output / "diesel_assets.csv",
                   ["grid_number", "grid_name", "unit_name", "dependable_capacity_mw",
                    "accumulated_operating_hours"], diesel_rows)

        manifest = {
            "status": "imported",
            "source_workbook": source.name,
            "site_count": len(sites),
            "row_counts": {
                "sites.csv": len(sites), "monthly_power_demand.csv": len(demand_rows),
                "monthly_ghi.csv": len(ghi_rows), "monthly_wind_speed.csv": len(wind_rows),
                "diesel_assets.csv": len(diesel_rows),
            },
            "source_sheets": list(REQUIRED_SHEETS),
            "data_limitations": [
                "Monthly power-demand values are summaries, not hourly observations; do not train an hourly forecast directly on them.",
                "The GHI unit is preserved as written and must be confirmed against original documentation before physical calculations.",
                "Wind speed is reported at 10 m elevation; hub-height conversion needs assumptions or measurements.",
                "Scenario sheets are excluded because scenario projections are not direct operational observations.",
                "The original source, measurement method, and licence are not independently verified by this importer.",
                "The workbook has no timestamped freshwater demand, desalination production, or reservoir inventory for water-model training.",
            ],
            "recommended_use": [
                "Cross-island screening and monthly scenario calibration.",
                "Comparing site-level load, solar/wind resource, and diesel-asset differences.",
                "Not a substitute for hourly meter data or water-system operating logs.",
            ],
        }
        (output / "manifest.json").write_text(
            json.dumps(manifest, indent=2, ensure_ascii=False) + "\n", encoding="utf-8"
        )
        return manifest
    finally:
        workbook.close()


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbook", help="Path to the Philippines DIB Excel workbook")
    parser.add_argument("--output-dir", required=True, help="Output directory for CSVs and manifest")
    args = parser.parse_args()
    print(json.dumps(import_workbook(args.workbook, args.output_dir), indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()
